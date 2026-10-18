"""Chuyển report/report.md sang report/report.docx (tiêu đề, đoạn, danh sách, bảng, code, hình).

Hình: lần đầu một đoạn nhắc tới `fig_xxx.png`, chèn ảnh results/figures/fig_xxx.png ngay sau đoạn đó kèm chú thích.
Dùng:  python scripts/report_to_docx.py
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "results" / "figures"
CAPTIONS = {
    "fig_tradeoff.png": "Đánh đổi Recall@10 – P95 ở 10 client, index đầy đủ (thanh: min–max qua 3 lần).",
    "fig_quality_vs_size.png": "Recall@10 theo kích thước index (mẫu 500 truy vấn test).",
    "fig_latency_vs_size.png": "P95 ở 10 client theo kích thước index.",
    "fig_latency_vs_clients.png": "P95 theo số client đồng thời, index đầy đủ.",
    "fig_build_time.png": "Thời gian build index theo kích thước corpus và builder.",
    "fig_build_memory.png": "RAM đỉnh khi build index.",
}
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_runs(par, text: str):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        else:
            par.add_run(part)


def main() -> None:
    lines = (ROOT / "report" / "report.md").read_text(encoding="utf-8").splitlines()
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.5)
    used_figs: set[str] = set()
    fig_no = 0
    i = 0
    para_buf: list[str] = []

    def flush_para():
        nonlocal fig_no
        if not para_buf:
            return
        text = " ".join(x.strip() for x in para_buf)
        para_buf.clear()
        p = doc.add_paragraph()
        add_runs(p, text)
        for f in re.findall(r"fig_[a-z_]+\.png", text):
            if f not in used_figs and (FIG / f).exists():
                used_figs.add(f)
                fig_no += 1
                doc.add_picture(str(FIG / f), width=Cm(14))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph(f"Hình {fig_no}. {CAPTIONS.get(f, f)}")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.runs[0].italic = True

    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            flush_para()
            j = i + 1
            code = []
            while j < len(lines) and not lines[j].startswith("```"):
                code.append(lines[j])
                j += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(code))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            i = j + 1
            continue
        if line.startswith("|"):
            flush_para()
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Light Grid Accent 1"
            for r_i, row in enumerate(rows):
                for c_i, cell in enumerate(row[:len(rows[0])]):
                    par = table.cell(r_i, c_i).paragraphs[0]
                    add_runs(par, cell)
                    for run in par.runs:
                        run.font.size = Pt(9.5)
                        if r_i == 0:
                            run.bold = True
            doc.add_paragraph()
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            flush_para()
            level = len(m.group(1))
            doc.add_heading(m.group(2), level=0 if level == 1 else level - 1)
        elif re.match(r"^\s*(-|\d+\.)\s+", line):
            flush_para()
            bullet = line.lstrip().startswith("-")
            text = re.sub(r"^\s*(-|\d+\.)\s+", "", line)
            j = i + 1
            while j < len(lines) and lines[j].startswith("  ") and not re.match(r"^\s*(-|\d+\.)\s+", lines[j]):
                text += " " + lines[j].strip()
                j += 1
            p = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_runs(p, text)
            i = j
            continue
        elif line.startswith(">"):
            flush_para()
            p = doc.add_paragraph()
            add_runs(p, line.lstrip("> ").strip())
            p.runs[0].italic = True if p.runs else None
        elif not line.strip():
            flush_para()
        else:
            para_buf.append(line)
        i += 1
    flush_para()
    out = ROOT / "report" / "report.docx"
    doc.save(out)
    print(out, f"{fig_no} hình")


if __name__ == "__main__":
    main()
